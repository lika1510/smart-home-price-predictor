import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()

    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BG = RGBColor(9, 13, 22)      # #090d16
    SURFACE_BG = RGBColor(17, 24, 39)  # #111827
    CYAN = RGBColor(56, 189, 248)      # #38bdf8
    EMERALD = RGBColor(52, 211, 153)   # #34d399
    PURPLE = RGBColor(192, 132, 252)   # #c084fc
    AMBER = RGBColor(251, 191, 36)     # #fbbf24
    TEXT_MAIN = RGBColor(248, 250, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)

    blank_layout = prs.slide_layouts[6]

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

    def add_header(slide, title_text, category_text="PROPINTEL AI PLATFORM"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = CYAN
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MAIN
        p2.font.name = "Calibri"

    # =========================================================================
    # SLIDE 1: Title & Overview
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1)

    # Title Card Box
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.1))
    card1.fill.solid()
    card1.fill.fore_color.rgb = SURFACE_BG
    card1.line.color.rgb = CYAN
    card1.line.width = Pt(1.5)

    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_top = Inches(0.6)
    tf1.margin_left = Inches(0.8)

    p = tf1.paragraphs[0]
    p.text = "SMART HOME PRICE & RISK PREDICTOR"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN

    p = tf1.add_paragraph()
    p.text = "R Language & Machine Learning-Powered Web Application (INR ₹)"
    p.font.size = Pt(20)
    p.font.color.rgb = CYAN
    p.space_before = Pt(10)

    p = tf1.add_paragraph()
    p.text = "\n• 4 Core Machine Learning Algorithms (2 Supervised + 2 Unsupervised)"
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_MUTED

    p = tf1.add_paragraph()
    p.text = "• Real-Time Property Valuation in Indian Rupees (₹ Lakhs / ₹ Cr)"
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_MUTED

    p = tf1.add_paragraph()
    p.text = "• React 18 Dynamic Frontend & Flask Integrated R Backend Engine"
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: Real-World Problem Statement & Solution
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "Real-World Problem & Solution Overview")

    # Column 1: Problem
    box_prob = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    box_prob.fill.solid()
    box_prob.fill.fore_color.rgb = SURFACE_BG
    box_prob.line.color.rgb = RGBColor(239, 68, 68)

    tf_p = box_prob.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = Inches(0.4)
    tf_p.margin_top = Inches(0.4)

    p = tf_p.paragraphs[0]
    p.text = "⚠️ THE PROBLEM"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(239, 68, 68)

    bullets_p = [
        "Home buyers & investors struggle to estimate fair market home valuations.",
        "Traditional appraisal models lack automated risk assessment tools.",
        "Manual classification of property market tiers is slow and prone to bias.",
        "Lack of real-time financing, EMI, and cashflow yield diagnostics."
    ]
    for b in bullets_p:
        p = tf_p.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(12)

    # Column 2: Solution
    box_sol = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    box_sol.fill.solid()
    box_sol.fill.fore_color.rgb = SURFACE_BG
    box_sol.line.color.rgb = EMERALD

    tf_s = box_sol.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = Inches(0.4)
    tf_s.margin_top = Inches(0.4)

    p = tf_s.paragraphs[0]
    p.text = "✅ THE SOLUTION"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    bullets_s = [
        "Automated Price Prediction in ₹ Lakhs / ₹ Cr using Linear Regression.",
        "Real-Time Deal Approval Risk Probability (%) using Logistic Regression.",
        "Unsupervised Property Tiering (Luxury Villa, Suburban Home, Fixer Plot).",
        "Interactive Home Loan EMI & Monthly Rental Cashflow Calculator."
    ]
    for b in bullets_s:
        p = tf_s.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(12)

    # =========================================================================
    # SLIDE 3: System Architecture & Data Flow
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "System Architecture & End-to-End Workflow")

    col_w = Inches(3.6)
    gap = Inches(0.4)
    left_start = Inches(0.8)

    # Step 1
    s1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_start, Inches(2.0), col_w, Inches(4.8))
    s1.fill.solid()
    s1.fill.fore_color.rgb = SURFACE_BG
    s1.line.color.rgb = CYAN

    tf = s1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.4)
    p = tf.paragraphs[0]
    p.text = "1. REACT FRONTEND UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p_b = ["React 18 Single Page App", "Tailwind CSS Styling", "Interactive Property Sliders", "Renovation What-If Simulator", "EMI Cashflow Calculator"]
    for b in p_b:
        p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(10)

    # Step 2
    s2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_start + col_w + gap, Inches(2.0), col_w, Inches(4.8))
    s2.fill.solid()
    s2.fill.fore_color.rgb = SURFACE_BG
    s2.line.color.rgb = EMERALD

    tf = s2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.4)
    p = tf.paragraphs[0]
    p.text = "2. FLASK BACKEND SERVER"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p_b = ["Flask REST Web App (app.py)", "Data Bridge (r_engine.py)", "Master Dataset Pipeline", "INR ₹ Currency Conversion", "API /api/predict Endpoint"]
    for b in p_b:
        p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(10)

    # Step 3
    s3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_start + (col_w + gap)*2, Inches(2.0), col_w, Inches(4.8))
    s3.fill.solid()
    s3.fill.fore_color.rgb = SURFACE_BG
    s3.line.color.rgb = PURPLE

    tf = s3.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.4)
    p = tf.paragraphs[0]
    p.text = "3. R ALGORITHMS ENGINE"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE
    p_b = ["Pure R Script (algorithms.R)", "lm() Linear Regression", "glm() Logistic Regression", "kmeans() Clustering", "prcomp() PCA Analysis"]
    for b in p_b:
        p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 4: The 4 Machine Learning Algorithms
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_bg(slide4)
    add_header(slide4, "The 4 Machine Learning Algorithms (Under The Hood)")

    algos = [
        ("1. Linear Regression", "Supervised Learning • lm()", "Predicts fair market home price in ₹ Lakhs / ₹ Cr and price/sqft based on SqFt, Bedrooms, Age, and Location.", CYAN),
        ("2. Logistic Regression", "Supervised Learning • glm()", "Predicts deal approval probability (0-100%) and acquisition risk rating (APPROVED vs REJECTED).", EMERALD),
        ("3. K-Means Clustering", "Unsupervised Learning • kmeans()", "Classifies properties into 3 market tiers (Luxury Villa, Suburban Home, Budget Fixer Plot).", PURPLE),
        ("4. PCA Analysis", "Unsupervised Learning • prcomp()", "Compresses multiple location metrics into a single composite Neighborhood Quality Score (0-100).", AMBER)
    ]

    card_w = Inches(5.6)
    card_h = Inches(2.3)

    positions = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]

    for idx, (title, sub, desc, color) in enumerate(algos):
        x, y = positions[idx]
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE_BG
        card.line.color.rgb = color

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 5: Key Results & Business Value
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_bg(slide5)
    add_header(slide5, "Key Results, Advantages & Business Value")

    res_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.0))
    res_box.fill.solid()
    res_box.fill.fore_color.rgb = SURFACE_BG
    res_box.line.color.rgb = CYAN

    tf_r = res_box.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.6)
    tf_r.margin_top = Inches(0.4)

    p = tf_r.paragraphs[0]
    p.text = "🎯 PROJECT IMPACT & RESULTS"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CYAN

    bullets_r = [
        "100% Automated Valuation: Delivers instant home prices in Indian Rupees (₹) with 95% Confidence Bounds.",
        "Integrated Risk Assessment: Eliminates bad real estate deals using logistic regression probability scoring.",
        "Interactive Financial Planning: Calculates home loan EMI (₹/mo), monthly rental income, and net cashflows.",
        "Simplified Architecture: Clean 4-file codebase (app.py, r_engine.py, algorithms.R, main.js) easy to maintain and scale."
    ]
    for b in bullets_r:
        p = tf_r.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(16)

    # Save presentation
    output_path = "/Users/vaddelikhitha/Desktop/lalala/Smart_Home_Price_Predictor_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    build_presentation()
